"""프로젝트 내보내기 라우터."""

import io
import json
import logging
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func

from database.connection import get_db
from database.models import Project, Component, CircuitBlock, DesignHistory, ChatMessage
from database.schemas import ComponentOut, CircuitBlockOut, DesignHistoryOut

logger = logging.getLogger(__name__)

router = APIRouter()


@router.post("/bom/{project_id}")
async def export_bom(
    project_id: int,
    db: AsyncSession = Depends(get_db),
):
    """프로젝트의 BoM을 Excel 파일로 내보냅니다."""
    # 프로젝트 존재 확인
    proj_result = await db.execute(
        select(Project).where(Project.id == project_id)
    )
    project = proj_result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="프로젝트를 찾을 수 없습니다.")

    # 부품 조회
    comp_result = await db.execute(
        select(Component)
        .where(Component.project_id == project_id)
        .order_by(Component.ref_designator)
    )
    components = comp_result.scalars().all()

    if not components:
        raise HTTPException(
            status_code=404,
            detail="내보낼 부품이 없습니다. 먼저 부품을 추가하세요.",
        )

    # Excel 생성
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        raise HTTPException(
            status_code=501,
            detail="openpyxl 패키지가 설치되지 않았습니다. 'pip install openpyxl'을 실행하세요.",
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"

    # 헤더 스타일
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # 프로젝트 정보 헤더
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = f"BOM - {project.name}"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:J2")
    desc_cell = ws["A2"]
    desc_cell.value = f"생성일: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} | 상태: {project.status}"
    desc_cell.alignment = Alignment(horizontal="center")

    # 컬럼 헤더 (4행)
    headers = [
        "No.", "Ref Designator", "Part Number", "Manufacturer",
        "Description", "Package", "Quantity", "Unit Price",
        "Total Price", "LCSC Code", "Category",
    ]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # 데이터 행
    total_cost = 0.0
    for row_idx, comp in enumerate(components, start=5):
        line_total = comp.unit_price * comp.quantity
        total_cost += line_total
        row_data = [
            row_idx - 4,
            comp.ref_designator,
            comp.part_number,
            comp.manufacturer,
            comp.description,
            comp.package,
            comp.quantity,
            comp.unit_price,
            round(line_total, 4),
            comp.lcsc_code,
            comp.category,
        ]
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            if col_idx in (8, 9):  # 가격 컬럼
                cell.number_format = "#,##0.0000"

    # 합계 행
    summary_row = len(components) + 5
    ws.cell(row=summary_row, column=6, value="합계:").font = Font(bold=True)
    ws.cell(row=summary_row, column=7, value=sum(c.quantity for c in components)).font = Font(bold=True)
    total_cell = ws.cell(row=summary_row, column=9, value=round(total_cost, 4))
    total_cell.font = Font(bold=True)
    total_cell.number_format = "#,##0.0000"

    # 컬럼 너비 조정
    col_widths = [6, 16, 20, 20, 35, 12, 10, 12, 12, 14, 16]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    # 메모리에 저장
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    safe_name = project.name.replace(" ", "_").replace("/", "_")
    filename = f"BOM_{safe_name}_{datetime.now().strftime('%Y%m%d')}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/report/{project_id}")
async def export_report(
    project_id: int,
    db: AsyncSession = Depends(get_db),
):
    """프로젝트 리포트를 JSON으로 내보냅니다."""
    # 프로젝트 조회
    proj_result = await db.execute(
        select(Project).where(Project.id == project_id)
    )
    project = proj_result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="프로젝트를 찾을 수 없습니다.")

    # 부품 조회
    comp_result = await db.execute(
        select(Component)
        .where(Component.project_id == project_id)
        .order_by(Component.ref_designator)
    )
    components = comp_result.scalars().all()

    # 회로 블록 조회
    block_result = await db.execute(
        select(CircuitBlock)
        .where(CircuitBlock.project_id == project_id)
        .order_by(CircuitBlock.name)
    )
    blocks = block_result.scalars().all()

    # 설계 이력 조회
    history_result = await db.execute(
        select(DesignHistory)
        .where(DesignHistory.project_id == project_id)
        .order_by(DesignHistory.created_at.desc())
        .limit(100)
    )
    history_entries = history_result.scalars().all()

    # 통계 계산
    total_cost = sum(c.unit_price * c.quantity for c in components)
    total_quantity = sum(c.quantity for c in components)
    categories = {}
    for c in components:
        cat = c.category or "미분류"
        if cat not in categories:
            categories[cat] = {"count": 0, "total_cost": 0.0}
        categories[cat]["count"] += c.quantity
        categories[cat]["total_cost"] += c.unit_price * c.quantity

    # JSON 리포트 구성
    report = {
        "report_type": "project_report",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "project": {
            "id": project.id,
            "name": project.name,
            "description": project.description,
            "spec": project.spec,
            "status": project.status,
            "created_at": project.created_at.isoformat() if project.created_at else None,
            "updated_at": project.updated_at.isoformat() if project.updated_at else None,
        },
        "statistics": {
            "component_count": len(components),
            "total_quantity": total_quantity,
            "total_estimated_cost": round(total_cost, 4),
            "circuit_block_count": len(blocks),
            "design_history_count": len(history_entries),
            "categories": categories,
        },
        "components": [
            ComponentOut.model_validate(c).model_dump() for c in components
        ],
        "circuit_blocks": [
            CircuitBlockOut.model_validate(b).model_dump() for b in blocks
        ],
        "design_history": [
            DesignHistoryOut.model_validate(h).model_dump() for h in history_entries
        ],
    }

    # JSON 스트리밍 응답
    report_json = json.dumps(report, ensure_ascii=False, indent=2, default=str)
    output = io.BytesIO(report_json.encode("utf-8"))

    safe_name = project.name.replace(" ", "_").replace("/", "_")
    filename = f"Report_{safe_name}_{datetime.now().strftime('%Y%m%d')}.json"

    return StreamingResponse(
        output,
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/formats")
async def list_export_formats():
    """지원하는 내보내기 형식 목록을 반환합니다."""
    formats = [
        {
            "format": "bom_excel",
            "description": "BoM Excel 파일 (.xlsx)",
            "endpoint": "/api/export/bom/{project_id}",
            "method": "POST",
        },
        {
            "format": "project_report_json",
            "description": "프로젝트 리포트 JSON (.json)",
            "endpoint": "/api/export/report/{project_id}",
            "method": "POST",
        },
    ]
    return {"formats": formats, "total": len(formats)}
